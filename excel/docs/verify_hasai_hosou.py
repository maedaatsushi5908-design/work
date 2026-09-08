#!/usr/bin/env python3
"""M_HasaiHosou が書き込む数式を、VBA と同じ手順で再現する。

作成環境に Excel が無いので、マクロの判定をそのまま Python に写して
「どのセルに何が入るか」を出す。VBA を直したらこちらも直すこと。

    python3 excel/docs/verify_hasai_hosou.py
"""
import os

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
WORK = os.environ.get("WORK", "01_higashishirakawa")
FOLDER = os.path.join(HERE, "..", "works", WORK)
BOOK = os.environ.get("BOOK", "06_dokou_hosou.xlsx")

# ---- M_HasaiHosou の先頭にある設定と同じもの -------------------------------
TARGET_SHEET = "総括表（舗装工事）"
PAVE_SRC = "舗装（集計）"
# I44・I46・I48は「殻運搬（現場→処分地）」の15cm以下（As車道・As歩道・
# Co舗装）、I45・I49は同じ枠の15cm超（As車道・Co舗装）。As歩道の15cm超
# （I47）は舗装（集計）側に対応するセルが無く、Co取壊（I50・I51）も
# 対応する枠が見つからないため見送っている。
CELL_MAP = "I7=M4|I8=M5|I9=P4|I10=P5|I44=M23|I45=M24|I46=M25|I48=P23|I49=P24"

# 舗装版破砕（舗装工事）機械の行 → 「As/Co の別:厚さの基準行[:追加で
# 拾う厚さ]」の対応。I24 だけは11行目（人力・As・4cm以下の厚さの定義行）
# を基準にする（11行目も24行目もH列は同じ「4」なので結果は変わらない）。
# 34行目（機械・Co・15cm以下）はCo破砕だけ15cmと14cmの区分があり、
# 14cmの分は15cmの欄にまとめている（As側は14cmが別行＝27行目に
# そのまま残るので、このまとめは無い）ため、No.2側（AF/AG列）から
# 厚さ14のぶんも追加で拾う。
HASAI_KIKAI_MAP = ("24=As:11|25=As:25|26=As:26|27=As:27|28=As:28|32=As:32|"
                    "33=Co:33|34=Co:34:14|35=Co:35")
HASAI_AS_NO1_THK = "$L$12:$L$19"
HASAI_AS_NO1_SUM = "$M$12:$M$19"
HASAI_AS_NO2_THK = "$AC$10:$AC$16"
HASAI_AS_NO2_SUM = "$AD$10:$AD$16"
HASAI_CO_NO1_THK = "$O$12:$O$19"
HASAI_CO_NO1_SUM = "$P$12:$P$19"
HASAI_CO_NO2_THK = "$AF$10:$AF$16"
HASAI_CO_NO2_SUM = "$AG$10:$AG$16"

INPUT_COLOR = "FFFF00"


def sheet_ref(sn):
    """数式に書くシート名。囲む必要のある名前だけ ' で囲む（VBA の SheetRef）"""
    def safe(ch):
        o = ord(ch)
        if ch.isascii() and (ch.isalnum() or ch == "_"):
            return True
        if o in (0x3000, 0x30FB):
            return False
        return (0x3041 <= o <= 0x30FF) or (0x4E00 <= o <= 0x9FFF) or (0xFF66 <= o <= 0xFF9F)

    if not sn or sn[0].isdigit() or not all(safe(c) for c in sn):
        return "'" + sn.replace("'", "''") + "'!"
    return sn + "!"


def is_input_cell(ws, addr):
    """黄色く塗ってある入力セルか（VBA の Interior.Color = 65535 と同じ）"""
    f = ws[addr].fill
    if f is None or f.patternType is None:
        return False
    fg = f.fgColor
    if fg.type == "rgb":
        return isinstance(fg.rgb, str) and fg.rgb.endswith(INPUT_COLOR)
    if fg.type == "indexed":
        return fg.indexed == 13        # 既定パレットの黄色 = FFFF00
    return False


def hasai_kikai_ref(sheet_name, r):
    """舗装版破砕（舗装工事）機械の1セル分（VBA の HasaiKikaiRef）"""
    spec = None
    for p in HASAI_KIKAI_MAP.split("|"):
        k, v = p.split("=")
        if int(k) == r:
            spec = v
            break
    if spec is None:
        return ""
    parts = spec.split(":")
    kind, h_row = parts[0], int(parts[1])
    if kind == "As":
        thk1, sum1, thk2, sum2 = (HASAI_AS_NO1_THK, HASAI_AS_NO1_SUM,
                                   HASAI_AS_NO2_THK, HASAI_AS_NO2_SUM)
    elif kind == "Co":
        thk1, sum1, thk2, sum2 = (HASAI_CO_NO1_THK, HASAI_CO_NO1_SUM,
                                   HASAI_CO_NO2_THK, HASAI_CO_NO2_SUM)
    else:
        return ""
    h_ref = sheet_ref(sheet_name) + "$H" + str(h_row)
    f = ("=SUMIF(" + sheet_ref(PAVE_SRC) + thk1 + "," + h_ref + "," +
         sheet_ref(PAVE_SRC) + sum1 + ")+SUMIF(" + sheet_ref(PAVE_SRC) + thk2 + "," +
         h_ref + "," + sheet_ref(PAVE_SRC) + sum2 + ")")

    # 3つ目の欄（追加の厚さ）があれば、No.2側からその厚さのぶんも
    # 追加で拾う（例：Co破砕の14cmを15cmの欄にまとめる場合）
    if len(parts) >= 3:
        for e in parts[2].split(","):
            f += ("+SUMIF(" + sheet_ref(PAVE_SRC) + thk2 + "," + e + "," +
                  sheet_ref(PAVE_SRC) + sum2 + ")")
    return f


def main():
    path = os.path.join(FOLDER, BOOK)
    wb = openpyxl.load_workbook(path, data_only=False)
    if TARGET_SHEET not in wb.sheetnames:
        print(f"シートが見つかりません: {TARGET_SHEET}")
        return 1
    if PAVE_SRC not in wb.sheetnames:
        print(f"転記元シートが見つかりません: {PAVE_SRC}")
        return 1
    ws = wb[TARGET_SHEET]

    written = {}
    skipped = []
    for p in CELL_MAP.split("|"):
        addr, cell_ref = p.split("=")
        if is_input_cell(ws, addr):
            written[addr] = "=" + sheet_ref(PAVE_SRC) + cell_ref
        else:
            skipped.append(addr)

    for p in HASAI_KIKAI_MAP.split("|"):
        r = int(p.split("=")[0])
        addr = "I" + str(r)
        if is_input_cell(ws, addr):
            f = hasai_kikai_ref(TARGET_SHEET, r)
            if f:
                written[addr] = f
        else:
            skipped.append(addr)

    ok = True
    for addr, expect in (
            ("I7", "='舗装（集計）'!M4"),
            ("I8", "='舗装（集計）'!M5"),
            ("I9", "='舗装（集計）'!P4"),
            ("I10", "='舗装（集計）'!P5"),
            ("I24", "=SUMIF('舗装（集計）'!$L$12:$L$19,'総括表（舗装工事）'!$H11,"
                     "'舗装（集計）'!$M$12:$M$19)"
                     "+SUMIF('舗装（集計）'!$AC$10:$AC$16,'総括表（舗装工事）'!$H11,"
                     "'舗装（集計）'!$AD$10:$AD$16)"),
            ("I34", "=SUMIF('舗装（集計）'!$O$12:$O$19,'総括表（舗装工事）'!$H34,"
                     "'舗装（集計）'!$P$12:$P$19)"
                     "+SUMIF('舗装（集計）'!$AF$10:$AF$16,'総括表（舗装工事）'!$H34,"
                     "'舗装（集計）'!$AG$10:$AG$16)"
                     "+SUMIF('舗装（集計）'!$AF$10:$AF$16,14,'舗装（集計）'!$AG$10:$AG$16)"),
            ("I44", "='舗装（集計）'!M23"), ("I45", "='舗装（集計）'!M24"),
            ("I46", "='舗装（集計）'!M25"), ("I48", "='舗装（集計）'!P23"),
            ("I49", "='舗装（集計）'!P24")):
        g = written.get(addr, "")
        mark = "一致" if g == expect else f"違う（{g}）"
        print(f"{addr} = {expect}  … {mark}")
        ok = ok and g == expect

    print("\n指示された式がすべて一致したか:", "はい" if ok else "いいえ")

    print(f"\n=== 書き込むセル {len(written)} 個 ===")
    for addr in sorted(written, key=lambda a: int(a[1:])):
        print(f"  {addr}: {written[addr]}")

    if skipped:
        print(f"\n見送り {len(skipped)} 個（黄色でないセル）: {', '.join(skipped)}")

    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
